"""Build plate v124 from v123: row 88 leaves both our courts.

One cell rewritten and two header lines replaced. Everything else is copied,
so a row nobody worked on keeps the words it had.

The frontend corrected its own previous letter: nothing in their code reads a
roster list at all. Three measurements on my side agree, and the log carries
no call that could have produced it. So 88 goes back to the tester with one
question rather than sitting amber in a court neither of us owns.
"""

import json
import sys

from docx import Document

SRC = "FOR_TORNIKE_MISHO_PLATE_2026-09-21_v123_SRULI.docx"
OUT = "FOR_TORNIKE_MISHO_PLATE_2026-09-21_v124_SRULI.docx"
ID_COLUMN = 1
COMMENT_COLUMN = 6


def set_cell(cell, text: str) -> None:
    """Rewrite a cell line by line, keeping the first run's formatting."""
    first = cell.paragraphs[0]
    template = first.runs[0] if first.runs else None
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)
    for run in list(first.runs):
        run._element.getparent().remove(run._element)

    def add(paragraph, line: str) -> None:
        run = paragraph.add_run(line)
        if template is not None:
            run.font.size = template.font.size
            run.font.name = template.font.name
            run.bold = template.bold

    lines = text.split("\n")
    add(first, lines[0])
    for line in lines[1:]:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_after = first.paragraph_format.space_after
        add(paragraph, line)


def set_paragraph(paragraph, text: str) -> None:
    template = paragraph.runs[0] if paragraph.runs else None
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    if template is not None:
        run.font.size = template.font.size
        run.font.name = template.font.name
        run.bold = template.bold


def main(spec_path: str) -> None:
    spec = json.load(open(spec_path, encoding="utf-8"))
    document = Document(SRC)
    table = document.tables[0]

    written = []
    for row in table.rows:
        row_id = row.cells[ID_COLUMN].text.strip()
        if row_id in spec["cells"]:
            set_cell(row.cells[COMMENT_COLUMN], spec["cells"][row_id])
            written.append(row_id)

    missing = sorted(set(spec["cells"]) - set(written))
    if missing:
        raise SystemExit(f"plate_v124: rows not found in the table: {missing}")

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("v123 —"):
            set_paragraph(paragraph, spec["intro"][0])
        elif paragraph.text.startswith("v122-თან შედარებით"):
            set_paragraph(paragraph, spec["intro"][1])

    document.save(OUT)
    print(f"{OUT}: rewrote {', '.join(written)}")


if __name__ == "__main__":
    main(sys.argv[1])
